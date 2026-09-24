#!/usr/bin/env python3
"""Three-slide Action Pro × Leafy Brew partnership deck."""

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.abspath(__file__))
CREAM = RGBColor(0xF6, 0xF1, 0xE8)
INK = RGBColor(0x1C, 0x19, 0x17)
GREEN = RGBColor(0x1E, 0x3A, 0x30)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x5C, 0x55, 0x4E)


def bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def text(slide, x, y, w, h, value, size=18, color=INK, font="Arial", bold=False, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = value
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = font
    run.font.bold = bold
    run.font.italic = italic
    return box


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    s1 = prs.slides.add_slide(blank)
    bg(s1, CREAM)
    text(s1, 0.55, 0.35, 6, 0.3, "PARTNERSHIP", size=12, color=RGBColor(0x6B, 0x4A, 0x32))
    text(s1, 0.5, 0.6, 12, 0.8, "Action Pro  ×  Leafy Brew", size=40, font="Georgia")
    text(s1, 0.55, 1.5, 12, 0.7,
         "Buy from one shop today. Show that receipt at the other shop. Get a discount. Both businesses grow.",
         size=18)
    s1.shapes.add_picture(os.path.join(ROOT, "photos/leafy-brew.jpg"), Inches(0.55), Inches(2.35), Inches(6.0), Inches(4.15))
    s1.shapes.add_picture(os.path.join(ROOT, "photos/action-pro-interior.jpg"), Inches(6.8), Inches(2.35), Inches(6.0), Inches(4.15))
    text(s1, 0.55, 6.6, 5.5, 0.55, "Leafy Brew\nCoffee & Positivitea Company", size=14, bold=True)
    text(s1, 6.8, 6.6, 5.5, 0.55, "Action Pro\nLifestyle and fashion store", size=14, bold=True)

    s2 = prs.slides.add_slide(blank)
    bg(s2, CREAM)
    text(s2, 0.55, 0.32, 6, 0.28, "HOW IT WORKS", size=12, color=RGBColor(0x6B, 0x4A, 0x32))
    text(s2, 0.5, 0.55, 12, 0.7, "Same day. Show the receipt.", size=36, font="Georgia")

    for x, photo, title, step1, step2, save in (
        (0.5, "photos/leafy-brew.jpg", "Buy at Leafy Brew",
         "1   Buy something at Leafy Brew.",
         "2   Show today’s invoice at Action Pro.",
         "Get X% off at Action Pro"),
        (7.05, None, "Buy at Action Pro",
         "1   Buy something at Action Pro.",
         "2   Show today’s invoice at Leafy Brew.",
         "Get X% off at Leafy Brew"),
    ):
        if photo:
            s2.shapes.add_picture(os.path.join(ROOT, photo), Inches(x), Inches(1.5), Inches(5.75), Inches(2.15))
        else:
            for i, name in enumerate(("action-pro-floor.jpg", "action-pro-makeup.jpg", "action-pro-racks.jpg")):
                s2.shapes.add_picture(os.path.join(ROOT, "photos", name), Inches(x + i * 1.95), Inches(1.5), Inches(1.88), Inches(2.15))
        text(s2, x, 3.75, 5.6, 0.45, title, size=26, font="Georgia")
        text(s2, x, 4.3, 5.6, 0.35, step1, size=16)
        text(s2, x, 4.7, 5.6, 0.35, step2, size=16)
        bar = s2.shapes.add_shape(1, Inches(x), Inches(5.2), Inches(5.75), Inches(0.62))
        bar.fill.solid()
        bar.fill.fore_color.rgb = GREEN
        bar.line.fill.background()
        text(s2, x + 0.2, 5.32, 5.3, 0.4, save, size=18, color=WHITE, bold=True)

    text(s2, 0.55, 6.15, 12, 0.5,
         "The receipt has to be from the same day. That is the whole partnership.", size=18)
    text(s2, 0.55, 6.9, 4, 0.3, "Page 1  ·  Page 2  ·  Page 3", size=12, color=MUTED)

    s3 = prs.slides.add_slide(blank)
    bg(s3, CREAM)
    text(s3, 0.55, 0.32, 8, 0.28, "FOR LEAFY BREW", size=12, color=RGBColor(0x6B, 0x4A, 0x32))
    text(s3, 0.5, 0.55, 12, 0.7, "Your customers get more than coffee.", size=36, font="Georgia")
    text(s3, 0.55, 1.35, 7.2, 0.7,
         "A Leafy Brew customer can take today’s receipt to Action Pro, get a discount, and shop.",
         size=16)
    points = [
        ("A discount", "They buy at Leafy Brew, show the same-day receipt at Action Pro, and get X% off."),
        ("A shopping experience", "Action Pro is a lifestyle and fashion store. Clothes, style, and everyday pieces, on the same day."),
        ("A stronger reason to visit", "Leafy Brew offers more than the cup. Guests leave with a discount and a place to shop."),
    ]
    y = 2.2
    for title, body in points:
        card = s3.shapes.add_shape(1, Inches(0.55), Inches(y), Inches(7.15), Inches(1.35))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.fill.background()
        text(s3, 0.8, y + 0.18, 6.7, 0.4, title, size=22, font="Georgia")
        text(s3, 0.8, y + 0.62, 6.7, 0.55, body, size=14, color=MUTED)
        y += 1.5
    s3.shapes.add_picture(os.path.join(ROOT, "photos/action-pro-floor.jpg"), Inches(8.0), Inches(2.2), Inches(4.75), Inches(2.35))
    s3.shapes.add_picture(os.path.join(ROOT, "photos/action-pro-makeup.jpg"), Inches(8.0), Inches(4.65), Inches(2.3), Inches(1.7))
    s3.shapes.add_picture(os.path.join(ROOT, "photos/action-pro-racks.jpg"), Inches(10.45), Inches(4.65), Inches(2.3), Inches(1.7))
    text(s3, 8.0, 6.5, 4.75, 0.55, "Action Pro\nLifestyle and fashion store", size=14, bold=True)

    out = os.path.join(ROOT, "Action-Pro-x-Leafy-Brew-Partnership.pptx")
    prs.save(out)
    print(out)


if __name__ == "__main__":
    main()
